"""Filling the college's DOCX template. No network, no LLM, no key.

The renderer is deliberately separable from generation: it takes finished text and
places it. That is what lets these tests prove the document is correct without a
vendor answering.
"""

from __future__ import annotations

import glob
import os

import pytest
from docx import Document
from docx.oxml.ns import qn

from pap.archives.deliverable import (MAPA_TEMPLATE, TEMPLATE_DIR, Identity,
                                      add_body, fill_identity, render_deliverable)


def _cells(document) -> list[str]:
    out = []
    for table in document.tables:
        for tc in table._tbl.iter(qn("w:tc")):
            text = "".join(n.text or "" for n in tc.iter(qn("w:t"))).strip()
            if text:
                out.append(text)
    return out


def _texts(document) -> list[str]:
    return [p.text.strip() for p in document.paragraphs if p.text.strip()]


# -- the template itself ----------------------------------------------------
def test_the_template_ships_with_the_repository():
    """Phase 3 cannot render without it, and it is not reconstructible from code."""
    assert os.path.exists(MAPA_TEMPLATE)


def test_the_committed_template_carries_no_personal_data():
    """It went into git, so the RA and name are stripped — the labels remain."""
    cells = _cells(Document(MAPA_TEMPLATE))
    assert "Nome:" in cells and "R.A" in cells
    joined = " ".join(cells)
    assert "12345678" not in joined
    assert "Zaveruka" not in joined


# -- identity ---------------------------------------------------------------
def test_every_identity_field_is_filled_exactly_once():
    """Regression: the header is a merged 3-column table, so `row.cells` yields the
    same cell up to three times. A row/column fill wrote one field repeatedly and
    skipped another — `Valor da atividade` was the one that vanished."""
    document = Document(MAPA_TEMPLATE)
    written = fill_identity(document, Identity(
        nome="Fulano", ra="12345678-9", curso="Engenharia de Software",
        disciplina="TÓPICOS EM COMPUTAÇÃO II", valor="3,50", prazo="05/07/2026"))

    assert sorted(written) == ["curso", "disciplina", "nome", "prazo", "ra", "valor"]
    cells = _cells(document)
    assert "Valor da atividade: 3,50" in cells
    assert "R.A 12345678-9" in cells
    assert len(cells) == len(set(cells)), "a value was written into more than one cell"


def test_a_missing_field_leaves_the_label_bare_rather_than_inventing_one():
    """A wrong RA on a submitted MAPA is worse than an obviously empty one."""
    document = Document(MAPA_TEMPLATE)
    fill_identity(document, Identity(nome="Fulano"))
    assert "R.A" in _cells(document)
    assert not any(c.startswith("R.A ") and c != "R.A" for c in _cells(document))


@pytest.mark.parametrize("template", sorted(glob.glob(
    os.path.join(TEMPLATE_DIR, "*.docx"))), ids=os.path.basename)
def test_every_committed_template_can_be_filled(template):
    """Regression, and the reason this is parametrized over the directory rather
    than over a hard-coded list: six real documents from Studeo spell the same two
    fields six ways — `Nome:`/`Acadêmico:` and `R.A`/`R.A.`/`R.A:`/`R.A.:`.
    Literal-prefix matching passed on the one blank template it was written against
    and silently filled nothing on four of the other five, so the submission would
    have gone out with no name and no RA. A template added later is covered the
    moment it lands in the directory."""
    document = Document(template)
    written = fill_identity(document, Identity(
        nome="Fulano de Tal", ra="99999999-9", curso="Engenharia de Software",
        disciplina="X", valor="3,50", prazo="01/01/2026"))

    assert sorted(written) == ["curso", "disciplina", "nome", "prazo", "ra", "valor"]
    cells = _cells(document)
    # Not "no duplicate cells": two of these documents carry content tables whose
    # cells legitimately repeat ("RF1", "Alta"). What must hold is that each
    # identity value was written into exactly one cell — the merged-row bug wrote
    # one value three times.
    for value in ("Fulano de Tal", "99999999-9"):
        assert sum(value in c for c in cells) == 1, f"{value!r} landed in more than one cell"


# Every spelling observed across the six source documents. Kept synthetic on
# purpose: the completed submissions that revealed them carry a real name and RA
# and do not belong in git, so the regression must not depend on them being here.
@pytest.mark.parametrize("label", ["Nome:", "Acadêmico:", "Academico:", "NOME:"])
def test_every_observed_name_spelling_is_matched(label):
    document = Document(MAPA_TEMPLATE)
    _set_first_run(next(document.tables[0]._tbl.iter(qn("w:tc"))), label)
    assert "nome" in fill_identity(document, Identity(nome="Fulano"))
    assert f"{label} Fulano" in _cells(document)


@pytest.mark.parametrize("label", ["R.A", "R.A.", "R.A:", "R.A.:", "RA:"])
def test_every_observed_ra_spelling_is_matched(label):
    """Four of these came from four different real documents. A literal-prefix
    matcher recognised exactly one and left the rest bare."""
    document = Document(MAPA_TEMPLATE)
    cells = list(document.tables[0]._tbl.iter(qn("w:tc")))
    _set_first_run(cells[1], label)
    assert "ra" in fill_identity(document, Identity(ra="12345678-9"))
    assert f"{label} 12345678-9" in _cells(document)


def _set_first_run(tc, text: str) -> None:
    runs = list(tc.iter(qn("w:t")))
    runs[0].text = text
    for extra in runs[1:]:
        extra.text = ""


# -- body -------------------------------------------------------------------
def test_a_bullet_list_keeps_its_markers_when_the_template_has_no_list_style():
    """Regression: this template defines no `List Bullet` style, so styled-only
    rendering turned every bullet into an ordinary paragraph and the enumeration
    disappeared into run-on prose."""
    document = Document(MAPA_TEMPLATE)
    add_body(document, "- Primeiro ponto\n- Segundo ponto")
    body = _texts(document)
    assert any(t.startswith(("•", "Primeiro")) for t in body)
    assert "• Primeiro ponto" in body or any(
        p.style.name == "List Bullet" for p in document.paragraphs)


def test_the_authors_own_numbering_is_preserved_across_blocks():
    """Regression, caught on a live generation: the model wrote sections 1/2/3,
    each separated by a blank line, and a per-block counter renumbered every one
    of them to "1." — three differently-titled questions all numbered 1 in a
    graded document. The number that was written is always the right one."""
    document = Document(MAPA_TEMPLATE)
    add_body(document, "1. Primeira\n\nTexto.\n\n2. Segunda\n\nTexto.\n\n3. Terceira")
    numbered = [t for t in _texts(document) if t[:1].isdigit()]
    assert numbered == ["1. Primeira", "2. Segunda", "3. Terceira"]


def test_a_genuine_single_block_list_keeps_its_numbers_too():
    document = Document(MAPA_TEMPLATE)
    add_body(document, "1. um\n2. dois\n3. tres")
    assert [t for t in _texts(document) if t[:1].isdigit()] == ["1. um", "2. dois", "3. tres"]


def test_a_list_that_does_not_start_at_one_is_not_renormalised():
    """Continuation lists exist; forcing them to start at 1 corrupts them."""
    document = Document(MAPA_TEMPLATE)
    add_body(document, "4. quatro\n5. cinco")
    assert [t for t in _texts(document) if t[:1].isdigit()] == ["4. quatro", "5. cinco"]


def test_headings_become_headings():
    document = Document(MAPA_TEMPLATE)
    add_body(document, "## Introdução\n\nTexto do parágrafo.")
    styles = {p.text.strip(): p.style.name for p in document.paragraphs if p.text.strip()}
    assert styles["Introdução"].startswith("Heading")


def test_inline_bold_is_rendered_and_the_asterisks_do_not_survive():
    """A half-supported markdown dialect puts stray asterisks in a graded document."""
    document = Document(MAPA_TEMPLATE)
    add_body(document, "Analisa **tópicos em computação** aplicados.")
    paragraph = [p for p in document.paragraphs if "Analisa" in p.text][0]
    assert "*" not in paragraph.text
    assert any(run.bold and "tópicos" in run.text for run in paragraph.runs)


# -- the whole document -----------------------------------------------------
def test_rendering_preserves_the_page_setup_the_college_handed_out(tmp_path):
    """The renderer fills their file rather than building a new one. A rebuilt
    document silently loses the A4 page, the 1-inch margins and the header logo."""
    out = render_deliverable("## Introdução\n\nTexto.", Identity(nome="Fulano"),
                             str(tmp_path / "mapa.docx"))
    section = Document(out).sections[0]
    assert round(section.page_width / 914400, 2) == 8.27
    assert round(section.page_height / 914400, 2) == 11.69
    assert round(section.left_margin / 914400, 2) == 1.0


def test_the_header_logo_survives_rendering(tmp_path):
    out = render_deliverable("Texto.", Identity(nome="Fulano"), str(tmp_path / "m.docx"))
    header = Document(out).sections[0].header
    assert any("image" in rel.reltype for rel in header.part.rels.values())


def test_the_title_line_is_not_disturbed(tmp_path):
    out = render_deliverable("Texto.", Identity(nome="Fulano"), str(tmp_path / "m.docx"))
    assert _texts(Document(out))[0].startswith("MAPA – Material de Avaliação")


def test_a_missing_template_names_the_file_rather_than_raising_from_docx(tmp_path):
    with pytest.raises(FileNotFoundError, match="patterns/templates"):
        render_deliverable("x", Identity(), str(tmp_path / "m.docx"),
                           template_path=str(tmp_path / "nope.docx"))


# -- tables -----------------------------------------------------------------
def test_a_markdown_table_becomes_a_real_word_table():
    """Regression, caught on a live MAPA generation: the model produced the
    consolidated risk table the brief asked for and it landed as seven paragraphs
    of raw pipes — separator row `|---|---|` included — in a graded document.
    `patterns/mapa.md` §6 records that content tables carry real weight; one
    sample has five of them."""
    document = Document(MAPA_TEMPLATE)
    before = len(document.tables)
    add_body(document, "| Risco | Probabilidade |\n|---|---|\n| Escopo | Alta |\n| Custo | Média |")

    assert len(document.tables) == before + 1
    table = document.tables[-1]
    assert (len(table.rows), len(table.columns)) == (3, 2), "separator row must not become a row"
    assert table.rows[0].cells[0].text.strip() == "Risco"
    assert table.rows[2].cells[1].text.strip() == "Média"
    assert not any(p.text.strip().startswith("|") for p in document.paragraphs)


def test_the_table_header_is_bold_and_the_grid_is_drawn():
    """The template defines no `Table Grid` style — the same gap that flattened the
    bullet lists — so borders are written into the XML or the table reads as
    unaligned columns of text."""
    document = Document(MAPA_TEMPLATE)
    add_body(document, "| A | B |\n|---|---|\n| 1 | 2 |")
    table = document.tables[-1]
    assert all(r.bold for r in table.rows[0].cells[0].paragraphs[0].runs if r.text.strip())
    assert "tblBorders" in table._tbl.xml


def test_a_ragged_table_pads_instead_of_raising():
    """A model occasionally emits a short row. Losing the table over it would be
    worse than a blank cell."""
    document = Document(MAPA_TEMPLATE)
    add_body(document, "| A | B | C |\n|---|---|---|\n| 1 | 2 |")
    table = document.tables[-1]
    assert len(table.columns) == 3
    assert table.rows[1].cells[2].text.strip() == ""


def test_a_lone_pipe_line_is_not_mistaken_for_a_table():
    """One line is not a table; treating it as one would swallow ordinary prose."""
    document = Document(MAPA_TEMPLATE)
    add_body(document, "| isto nao e uma tabela |")
    assert len(document.tables) == 1  # identity only
